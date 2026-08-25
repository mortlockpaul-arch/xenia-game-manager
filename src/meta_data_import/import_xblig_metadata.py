from __future__ import annotations

import argparse
import re
import sqlite3
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


def normalize_title(title: str) -> str:
    """Normalise a title for matching against XBLIGGame.title."""
    value = str(title or "").casefold().strip()
    value = value.replace("&", "and")
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def text(value) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def iso_date(value) -> str | None:
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date().isoformat()

    value = text(value)
    if not value:
        return None

    for fmt in ("%m/%d/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            pass

    return value


def number(value) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def integer(value) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def header_map(ws) -> dict[str, int]:
    headers = {}
    for index, value in enumerate(
        next(ws.iter_rows(min_row=1, max_row=1, values_only=True)),
        start=1,
    ):
        if value is not None:
            headers[str(value).strip()] = index
    return headers


def cell(row, headers, name):
    index = headers.get(name)
    if index is None:
        return None
    return row[index - 1]


def ensure_schema(connection: sqlite3.Connection):
    connection.executescript(
        Path(__file__).with_name("xblig_metadata.sql").read_text(
            encoding="utf-8"
        )
    )


def import_xblig_games(
    database: Path,
    workbook: Path,
) -> tuple[int, int]:
    connection = sqlite3.connect(database)

    try:
        ensure_schema(connection)

        wb = load_workbook(
            workbook,
            read_only=True,
            data_only=True,
        )

        ws = wb["XBLIG Games"]
        headers = header_map(ws)

        title_column = (
            "XBLIG Marketplace Game Name "
            "(with archive.org link to old Marketplace page if found)"
        )

        if title_column not in headers:
            raise RuntimeError(
                f"Could not find title column: {title_column!r}"
            )

        inserted = 0
        updated = 0

        for row in ws.iter_rows(
            min_row=2,
            values_only=True,
        ):
            raw_title = cell(row, headers, title_column)
            title = text(raw_title)

            if not title:
                continue

            normalized = normalize_title(title)

            values = {
                "title": title,
                "normalized_title": normalized,
                "developer": text(
                    cell(row, headers, "Developer(s)")
                ),
                "developer_account": text(
                    cell(row, headers, "XBLIG Developer account name")
                ),
                "genre": text(
                    cell(row, headers, "XBLIG Genre")
                ),
                "release_date": iso_date(
                    cell(row, headers, "XBLIG Release Date MM/DD/YYYY")
                ),
                "user_rating": number(
                    cell(row, headers, "XBLIG UserRating (out of 5)")
                ),
                "rating_count": integer(
                    cell(row, headers, "XBLIG rating count")
                ),
                "file_size": text(
                    cell(row, headers, "XBLIG File Size")
                ),
                "available_on": text(
                    cell(row, headers, "Available on")
                ),
                "links": text(
                    cell(row, headers, "Links")
                ),
                "notes": text(
                    cell(row, headers, "Notes")
                ),
                "updates": text(
                    cell(row, headers, "Updates (as of 12/12/2010)")
                ),
                "category_scores": text(
                    cell(
                        row,
                        headers,
                        "Category Scores: Violence, Sex, Mature Content (out of 3)",
                    )
                ),
                "timecode": text(
                    cell(
                        row,
                        headers,
                        "Timecode in Youtube video https://youtu.be/oyH2a5dxjLE",
                    )
                ),
                "gamefaqs": 1
                if text(cell(row, headers, "On GameFAQs"))
                else 0,
                "title_text_from_youtube": text(
                    cell(
                        row,
                        headers,
                        "Title text from Youtube video ",
                    )
                ),
            }

            cursor = connection.execute(
                """
                SELECT id
                FROM xblig_metadata
                WHERE normalized_title = ?
                """,
                (normalized,),
            )

            exists = cursor.fetchone() is not None

            connection.execute(
                """
                INSERT INTO xblig_metadata (
                    title,
                    normalized_title,
                    developer,
                    developer_account,
                    genre,
                    release_date,
                    user_rating,
                    rating_count,
                    file_size,
                    available_on,
                    links,
                    notes,
                    updates,
                    category_scores,
                    timecode,
                    gamefaqs,
                    title_text_from_youtube,
                    source_file
                )
                VALUES (
                    :title,
                    :normalized_title,
                    :developer,
                    :developer_account,
                    :genre,
                    :release_date,
                    :user_rating,
                    :rating_count,
                    :file_size,
                    :available_on,
                    :links,
                    :notes,
                    :updates,
                    :category_scores,
                    :timecode,
                    :gamefaqs,
                    :title_text_from_youtube,
                    :source_file
                )
                ON CONFLICT(normalized_title) DO UPDATE SET
                    title = excluded.title,
                    developer = excluded.developer,
                    developer_account = excluded.developer_account,
                    genre = excluded.genre,
                    release_date = excluded.release_date,
                    user_rating = excluded.user_rating,
                    rating_count = excluded.rating_count,
                    file_size = excluded.file_size,
                    available_on = excluded.available_on,
                    links = excluded.links,
                    notes = excluded.notes,
                    updates = excluded.updates,
                    category_scores = excluded.category_scores,
                    timecode = excluded.timecode,
                    gamefaqs = excluded.gamefaqs,
                    title_text_from_youtube =
                        excluded.title_text_from_youtube,
                    source_file = excluded.source_file,
                    imported_at = CURRENT_TIMESTAMP
                """,
                {
                    **values,
                    "source_file": workbook.name,
                },
            )

            if exists:
                updated += 1
            else:
                inserted += 1

        connection.commit()
        return inserted, updated

    finally:
        connection.close()


def main():
    parser = argparse.ArgumentParser(
        description="Import the XBLIG XLSX catalogue into SQLite."
    )
    parser.add_argument(
        "database",
        type=Path,
        help="SQLite database file",
    )
    parser.add_argument(
        "workbook",
        type=Path,
        help="XBLIG XLSX workbook",
    )

    args = parser.parse_args()

    inserted, updated = import_xblig_games(
        args.database,
        args.workbook,
    )

    print(f"Inserted: {inserted:,}")
    print(f"Updated:  {updated:,}")


if __name__ == "__main__":

    inserted, updated = import_xblig_games(
        Path(r"C:\PycharmProjects\xenia-game-manager\src\db\games.db"),
        Path(r"/meta_data_import/List of Xbox Live Indie Games.xlsx"),
    )

    print(f"Inserted: {inserted:,}")
    print(f"Updated:  {updated:,}")
